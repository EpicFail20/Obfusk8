# Copyright (C) 2026 CARROLAGGI Xavier
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.
"""
The common normalization (text_normalization.py, D-014) is applied by every
flow at the point where text is sent to the analyzer: text API, DOCX, CSV,
PDF, image. The analyzer sees the normalized text; every position and every
redaction refers to the text RECEIVED, and covers the invisible characters
inside a value.

Presidio is replaced by a fake `main._analyze_text` that finds the plain
form "Camille Martin" (or "Chloé Boyer") and nothing else: before the
normalization, the variants below were invisible to it, as the phase 1
benchmark measured on the real analyzer.
"""

import io
import json
import re
import unicodedata

import pymupdf as fitz
import pytest
from docx import Document as WordDocument
from fastapi import FastAPI

import main
import text_api
from tests.doc_headers import doc_headers
from tests.test_main_units import _drive, _FakeRequest, _SyncUpload
from tests.test_text_api import EXT_ORIGIN, _post

NAME = "Camille Martin"
VARIANTS = {
    "largeur_nulle": "Camille\u200b Martin",
    "espace_insecable": "Camille\u00a0Martin",
    "espace_fine": "Camille\u202fMartin",
    "nfd": unicodedata.normalize("NFD", "Camille Martin"),
    "trait_conditionnel": "Cam\u00adille Martin",
    "bidi": "Camille \u202eMartin\u202c",
    "nul": "Camille\x00Martin",
}


def _fake_analyzer(seen: list[str] | None = None):
    def fake_analyze(text: str, theme: dict | None = None, **kwargs: object) -> list[dict]:
        if seen is not None:
            seen.append(text)
        return [
            {"entity_type": "PERSON", "start": m.start(), "end": m.end(), "score": 0.85}
            for m in re.finditer(r"Camille Martin|Chloé Boyer", text)
        ]

    return fake_analyze


@pytest.mark.parametrize("variant", VARIANTS.values(), ids=VARIANTS.keys())
def test_blocs_texte_positions_sur_le_texte_recu(monkeypatch, variant):
    """DOCX paragraphs and CSV cells go through _detect_text_blocks."""
    monkeypatch.setattr(main, "_analyze_text", _fake_analyzer())
    block = f"Le patient {variant} est sorti."
    detections = main._detect_text_blocks([(0, "Rien ici."), (1, block)])
    ner = [d for d in detections if d["source"] == "ner"]
    assert len(ner) == 1
    # A format character right after the value (closing bidi control) is
    # outside it: only what lies between its first and last visible
    # characters belongs to the value.
    assert block[ner[0]["start"] : ner[0]["end"]] == variant.rstrip("\u202c")


def test_normalisations_existantes_conservees(monkeypatch):
    """Upper-case words and typographic dashes are still normalized, after
    the new normalization."""
    seen: list[str] = []
    monkeypatch.setattr(main, "_analyze_text", _fake_analyzer(seen))
    main._detect_text_blocks([(0, "CAMILLE\u00a0MARTIN – 01–02–1990")])
    assert seen == ["Camille Martin - 01-02-1990"]


def test_texte_ordinaire_inchange(monkeypatch):
    """Nothing to normalize: the analyzer receives exactly what it received
    before (identity path)."""
    seen: list[str] = []
    monkeypatch.setattr(main, "_analyze_text", _fake_analyzer(seen))
    detections = main._detect_text_blocks([(0, "Bonjour Camille Martin"), (1, "Au revoir")])
    assert seen == ["Bonjour Camille Martin\nAu revoir"]
    assert [(d["start"], d["end"]) for d in detections if d["source"] == "ner"] == [(8, 22)]


def _docx(paragraphs: list[str]) -> bytes:
    document = WordDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _finalize(job_id: str) -> dict:
    response = _drive(
        main.finalize_document(
            _FakeRequest(doc_headers()),
            job_id=job_id,
            excluded_ids="",
            manual_zones="[]",
            redacted_image_ids="",
            response_format="json",
        )
    )
    assert response.status_code == 200
    return json.loads(bytes(response.body))


def _detect(filename: str, data: bytes) -> str:
    response = _drive(main.detect_document(_FakeRequest(doc_headers()), file=_SyncUpload(filename, data), theme=""))
    assert response.status_code == 200
    match = re.search(r'name="job_id" value="([0-9a-f]{32})"', bytes(response.body).decode("utf-8"))
    assert match
    return match.group(1)


@pytest.mark.parametrize("variant", ["Camille\u200b Martin", "Camille\u00a0Martin"])
def test_docx_caviardage_couvre_la_valeur_et_son_invisible(monkeypatch, variant):
    monkeypatch.setattr(main, "_analyze_text", _fake_analyzer())
    job_id = _detect("t.docx", _docx([f"Le patient {variant} est sorti."]))
    try:
        _finalize(job_id)
        out = main.WORKDIR / f"{job_id}-document-anonymise.docx"
        text = "\n".join(p.text for p in WordDocument(str(out)).paragraphs)
        out.unlink(missing_ok=True)
    finally:
        main.PENDING_JOBS.pop(job_id, None)
    assert text == f"Le patient {main.REDACTION_MARKER} est sorti."


def test_csv_caviardage_couvre_la_valeur_et_son_invisible(monkeypatch):
    monkeypatch.setattr(main, "_analyze_text", _fake_analyzer())
    # Neutral header: a recognized "nom" header would trigger the existing
    # structural rules (EXT-40), which are not what is tested here.
    job_id = _detect("t.csv", "ligne;texte\n1;Le patient Camille\u200b Martin est sorti\n".encode())
    try:
        _finalize(job_id)
        out = main.WORKDIR / f"{job_id}-document-anonymise.csv"
        text = out.read_text(encoding="utf-8")
        out.unlink(missing_ok=True)
    finally:
        main.PENDING_JOBS.pop(job_id, None)
    assert f"1;Le patient {main.REDACTION_MARKER} est sorti" in text
    assert "\u200b" not in text, "the invisible character inside the value is redacted with it"


def test_pdf_nul_entre_deux_mots_detecte_et_caviarde(monkeypatch):
    """EXT-38: "Chloé<NUL>Boyer" (glyph without Unicode mapping) is now seen
    as "Chloé Boyer" by the analyzer, located (D-026 fallback) and redacted."""
    monkeypatch.setattr(main, "_analyze_text", _fake_analyzer())
    doc = fitz.open()
    page = doc.new_page()
    page.insert_font(fontname="mono", fontfile="/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")
    page.insert_text((50, 80), "Suivi par le Dr Chloé\u200bBoyer demain.", fontname="mono", fontsize=10)
    pdf = doc.tobytes()
    doc.close()
    doc = fitz.open(stream=pdf, filetype="pdf")
    try:
        assert "Chloé\x00Boyer" in doc[0].get_text()
        detections, report = main._detect_pdf(doc)
        main._apply_selected_redactions(doc, detections, set())
        remaining = doc[0].get_text()
    finally:
        doc.close()
    assert "Boyer" not in remaining and "Chlo" not in remaining
    assert "Suivi" in remaining and "demain" in remaining
    assert not report["unlocated"] and not report["incomplete"]


def test_image_mot_ocr_avec_invisible(monkeypatch):
    """OCR words are joined into a text; the same normalization applies and
    the entity maps back to the word boxes."""
    monkeypatch.setattr(main, "_analyze_text", _fake_analyzer())
    words = [
        {
            "text": "Patient",
            "left": 0,
            "top": 0,
            "width": 50,
            "height": 10,
            "block_num": 1,
            "par_num": 1,
            "line_num": 1,
        },
        {
            "text": "Camille\u00a0Martin",
            "left": 60,
            "top": 0,
            "width": 90,
            "height": 10,
            "block_num": 1,
            "par_num": 1,
            "line_num": 1,
        },
        {
            "text": "sorti",
            "left": 160,
            "top": 0,
            "width": 30,
            "height": 10,
            "block_num": 1,
            "par_num": 1,
            "line_num": 1,
        },
    ]
    monkeypatch.setattr(main, "_run_ocr", lambda img: words)
    detections = main._detect_image(object())
    assert [d["page_rect"] for d in detections] == [[60.0, 0.0, 150.0, 10.0]]


@pytest.mark.parametrize("prefix", ["", "\U0001f600 ", "\U0001f1eb\U0001f1f7\U0001f600 "])
def test_api_texte_positions_et_utf16_sur_le_texte_recu(audit_dir, monkeypatch, prefix):
    monkeypatch.setattr(main, "_analyze_text", _fake_analyzer())
    monkeypatch.setattr(main, "AUDIT_DIR", audit_dir)
    app = FastAPI()
    app.include_router(
        main._build_text_api_router(text_api.TextApiSettings(enabled=True, allowed_origins=frozenset({EXT_ORIGIN})))
    )
    text = f"{prefix}Bonjour Camille\u200b Martin \U0001f600 et Cam\u00adille Martin."
    status, _, body, _ = _post(app, "/api/v1/text/analyze", {"text": text})
    assert status == 200, body
    persons = [e for e in body["entities"] if e["entity_type"] == "PERSON"]
    assert [text[e["start"] : e["end"]] for e in persons] == ["Camille\u200b Martin", "Cam\u00adille Martin"]
    utf16 = text.encode("utf-16-le")
    for e in persons:
        assert utf16[2 * e["start_utf16"] : 2 * e["end_utf16"]].decode("utf-16-le") == text[e["start"] : e["end"]]

    status, _, body, _ = _post(app, "/api/v1/text/pseudonymize", {"text": text})
    assert status == 200, body
    assert "Martin" not in body["text"] and "\u200b" not in body["text"] and "\u00ad" not in body["text"]


@pytest.fixture
def audit_dir(tmp_path):
    yield tmp_path
    import logging

    logger = logging.getLogger(text_api.AUDIT_LOGGER_NAME)
    for handler in list(logger.handlers):
        logger.removeHandler(handler)
        handler.close()
