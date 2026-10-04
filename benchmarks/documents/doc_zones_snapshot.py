# Copyright (C) 2026 CARROLAGGI Xavier
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.
"""
Document-flow zone snapshot (phase 2, non-regression of the redaction zones).

Phase 2 changes the document flow (Unicode normalization, recognizers for
all flows, worker thread). Its acceptance rule is "same redaction zones as
the reference, except additions listed one by one". This tool records the
zones proposed for review by /api/detect, THROUGH THE FULL CHAIN, for a fixed
set of synthetic documents, and compares two snapshots.

  snapshot:  python benchmarks/documents/doc_zones_snapshot.py
  compare:   python benchmarks/documents/doc_zones_snapshot.py --compare REF.json NEW.json
             (exit code 1 if a reference zone is lost)

Documents (rebuilt identically on every run, nothing stored): 24 prompts of
the quality corpus (the first 3 of each variant, excluding the `code`
category), rendered as a PDF with an embedded DejaVu font, a DOCX
(paragraphs and a table), a CSV and a PNG (OCR). Each is sent with no theme
and with each theme.

A zone is identified by its position and its entity types:
  - PDF/image: page and rectangle in preview pixels (rounded);
  - DOCX/CSV: rendered block (DOCX paragraph or CSV cell) number and
    character interval within it.
A reference zone missing from a new snapshot is "extended" when a current
zone still masks all of it, "lost" otherwise; only lost zones fail the
comparison.
The detected text itself is NEVER written to the report: only its length
and a truncated SHA-256 (the corpus contains synthetic secrets).

Run it in a throwaway client container of the app image (PyMuPDF,
python-docx and Pillow are there), as for e2e_document_flow.py:
  docker run --rm --network host -e BENCH_USER -e BENCH_PASSWORD -e BENCH_INSECURE_TLS=1 \\
    -e BENCH_RESOLVE_LOOPBACK=1 -v "$PWD:/repo" -w /repo \\
    --entrypoint python ghcr.io/epicfail20/obfusk8-app:main benchmarks/documents/doc_zones_snapshot.py
Each upload is paced (DOC_PACE seconds, 13 by default) under the existing
upload rate limit (5/min). Jobs are never finalized: nothing is written to
audit.log; they expire after JOB_REVIEW_TTL_SECONDS. The 16 jobs of one run
stay under MAX_PENDING_JOBS (20): wait for that TTL before running again.
"""

import argparse
import csv
import hashlib
import html
import io
import json
import os
import re
import sys
import textwrap
import time
from pathlib import Path
from typing import Any

import docx
import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from obfusk8_client import BASE_URL, HTTP_OK, login  # noqa: E402

CORPUS = HERE.parent / "quality" / "corpus.jsonl"
RESULTS = HERE.parent / "results"
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
PROMPTS_PER_VARIANT = 3
PROMPTS_PER_PDF_PAGE = 4
THEMES = ["", "medical", "it", "compta"]
PACE = float(os.environ.get("DOC_PACE", "13"))

_PAGE_START_RE = re.compile(r'<div class="page-container" data-page="(\d+)"')
_PIXEL_ZONE_RE = re.compile(
    r'<div class="detection" data-id="[0-9a-f]+" data-group="[0-9a-f]*" title="([^"]*)" '
    r'style="left:(-?\d+)px; top:(-?\d+)px; width:(-?\d+)px; height:(-?\d+)px;"'
)
# DOCX blocks are <div class="doc-block">; CSV cells are <td> (one block per cell).
_BLOCK_RE = re.compile(r'<div class="doc-block">(.*?)</div>|<td[^>]*>(.*?)</td>', re.S)
_TEXT_ZONE_RE = re.compile(
    r'<span class="text-detection" data-id="[0-9a-f]+" data-group="[0-9a-f]*" title="([^"]*)"[^>]*>(.*?)</span>', re.S
)
_LABEL_RE = re.compile(r'<span class="doc-block-label">.*?</span>', re.S)


def prompts() -> list[str]:
    """The first PROMPTS_PER_VARIANT prompts of each variant, `code` excluded."""
    taken: dict[str, int] = {}
    out = []
    for line in CORPUS.read_text(encoding="utf-8").splitlines():
        doc = json.loads(line)
        if doc["category"] == "code" or taken.get(doc["variant"], 0) >= PROMPTS_PER_VARIANT:
            continue
        taken[doc["variant"]] = taken.get(doc["variant"], 0) + 1
        out.append(doc["text"])
    return out


def make_pdf(texts: list[str]) -> bytes:
    doc = fitz.open()
    for i in range(0, len(texts), PROMPTS_PER_PDF_PAGE):
        page = doc.new_page()
        page.insert_font(fontname="dejavu", fontfile=FONT)
        y = 50.0
        for text in texts[i : i + PROMPTS_PER_PDF_PAGE]:
            rect = fitz.Rect(50, y, page.rect.width - 50, y + 180)
            page.insert_textbox(rect, text, fontname="dejavu", fontsize=10)
            y += 185
    out: bytes = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out


def make_docx(texts: list[str]) -> bytes:
    document = docx.Document()
    body, in_table = texts[:-4], texts[-4:]
    for text in body:
        document.add_paragraph(text)
    table = document.add_table(rows=2, cols=2)
    for index, text in enumerate(in_table):
        table.cell(index // 2, index % 2).text = text
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_csv(texts: list[str]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(["Ligne", "Texte"])
    for index, text in enumerate(texts, 1):
        writer.writerow([str(index), text])
    return buf.getvalue().encode("utf-8")


def make_png(texts: list[str]) -> bytes:
    """Only the first 8 prompts: OCR cost, and enough to cover the image path."""
    lines: list[str] = []
    for text in texts[:8]:
        for paragraph in text.split("\n"):
            lines.extend(textwrap.wrap(paragraph, 95) or [""])
        lines.append("")
    font = ImageFont.truetype(FONT, 22)
    img = Image.new("RGB", (1500, 40 + 32 * len(lines)), "white")
    draw = ImageDraw.Draw(img)
    for index, line in enumerate(lines):
        draw.text((20, 20 + 32 * index), line, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _digest(text: str) -> dict[str, Any]:
    return {"len": len(text), "sha": hashlib.sha256(text.encode("utf-8")).hexdigest()[:12]}


def pixel_zones(page_html: str) -> list[dict[str, Any]]:
    zones = []
    starts = list(_PAGE_START_RE.finditer(page_html))
    for index, match in enumerate(starts):
        end = starts[index + 1].start() if index + 1 < len(starts) else len(page_html)
        for types, left, top, width, height in _PIXEL_ZONE_RE.findall(page_html[match.end() : end]):
            zones.append(
                {
                    "page": int(match.group(1)),
                    "rect": [int(left), int(top), int(width), int(height)],
                    "types": html.unescape(types),
                }
            )
    return sorted(zones, key=lambda z: (z["page"], z["rect"], z["types"]))


def text_zones(page_html: str) -> list[dict[str, Any]]:
    zones = []
    for block_index, (docx_block, csv_cell) in enumerate(_BLOCK_RE.findall(page_html)):
        rest = _LABEL_RE.sub("", docx_block or csv_cell, count=1)
        pos = 0
        last = 0
        for match in _TEXT_ZONE_RE.finditer(rest):
            pos += len(html.unescape(rest[last : match.start()]))
            segment = html.unescape(match.group(2))
            zones.append(
                {
                    "block": block_index,
                    "start": pos,
                    "end": pos + len(segment),
                    "types": html.unescape(match.group(1)),
                    **_digest(segment),
                }
            )
            pos += len(segment)
            last = match.end()
    return zones


CASES = {
    "pdf": ("snapshot.pdf", make_pdf, pixel_zones),
    "docx": ("snapshot.docx", make_docx, text_zones),
    "csv": ("snapshot.csv", make_csv, text_zones),
    "image": ("snapshot.png", make_png, pixel_zones),
}


def snapshot() -> dict[str, Any]:
    texts = prompts()
    session = login()
    cases: dict[str, Any] = {}
    for theme in THEMES:
        for kind, (name, build, parse) in CASES.items():
            time.sleep(PACE)
            response = session.post(
                f"{BASE_URL}/api/detect", files={"file": (name, build(texts))}, data={"theme": theme}, timeout=180
            )
            if response.status_code != HTTP_OK or 'name="job_id"' not in response.text:
                raise RuntimeError(f"detect {kind} theme={theme or '-'}: HTTP {response.status_code}")
            zones = parse(response.text)
            cases[f"{kind}/{theme or 'aucun'}"] = zones
            print(f"{kind:6} theme={theme or '-':8} {len(zones)} zone(s)", flush=True)
    return {
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "prompts": len(texts),
        "corpus_sha256": hashlib.sha256(CORPUS.read_bytes()).hexdigest()[:16],
        "cases": cases,
    }


def _key(zone: dict[str, Any]) -> str:
    return json.dumps(zone, sort_keys=True, ensure_ascii=False)


# Rounding of the preview rectangles (whole pixels) between two snapshots.
PIXEL_TOLERANCE = 1


def _covers(outer: dict[str, Any], inner: dict[str, Any]) -> bool:
    """True when zone `outer` masks everything zone `inner` masked (same
    block or page), whatever their entity types."""
    if "block" in inner:
        return outer["block"] == inner["block"] and outer["start"] <= inner["start"] and inner["end"] <= outer["end"]
    if outer["page"] != inner["page"]:
        return False
    (ol, ot, ow, oh), (il, it, iw, ih) = outer["rect"], inner["rect"]
    tol = PIXEL_TOLERANCE
    return ol - tol <= il and ot - tol <= it and il + iw <= ol + ow + tol and it + ih <= ot + oh + tol


def compare(reference: dict[str, Any], current: dict[str, Any]) -> int:
    """Prints, per case, the reference zones that disappeared — "extended"
    when a current zone still masks all of it (a larger zone), "lost"
    otherwise — and the zones added. Exit code 1 if any zone is lost (the
    only regression the phase forbids)."""
    lost_total = 0
    for case in sorted(set(reference["cases"]) | set(current["cases"])):
        ref_zones, cur_zones = reference["cases"].get(case, []), current["cases"].get(case, [])
        before, after = {_key(z) for z in ref_zones}, {_key(z) for z in cur_zones}
        gone = [z for z in ref_zones if _key(z) not in after]
        extended = [z for z in gone if any(_covers(c, z) for c in cur_zones)]
        lost = [z for z in gone if not any(_covers(c, z) for c in cur_zones)]
        added = sorted(after - before)
        lost_total += len(lost)
        print(
            f"{case:16} reference={len(before)} current={len(after)} lost={len(lost)} "
            f"extended={len(extended)} added={len(added)}"
        )
        for label, items in (
            ("  - lost", [_key(z) for z in lost]),
            ("  ~ extended", [_key(z) for z in extended]),
            ("  + added", added),
        ):
            for item in items:
                print(f"{label} {item}")
    return 1 if lost_total else 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0] if __doc__ else None)
    parser.add_argument("--compare", nargs=2, metavar=("REF", "NEW"))
    args = parser.parse_args()
    if args.compare:
        ref, new = (json.loads(Path(p).read_text(encoding="utf-8")) for p in args.compare)
        sys.exit(compare(ref, new))
    report = snapshot()
    RESULTS.mkdir(exist_ok=True)
    out = RESULTS / f"doc-zones-{time.strftime('%Y%m%dT%H%M%S')}.json"
    out.write_text(json.dumps(report, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(out)
