# Copyright (C) 2026 CARROLAGGI Xavier
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.
"""
Document flow non-regression THROUGH THE FULL CHAIN (Traefik, oauth2-proxy,
Keycloak, app under the enforcing seccomp profile), synthetic content only:
detect -> preview -> finalize -> download for PDF, DOCX, CSV and image, with
and without theme; then the /api/v1/ routes compared to an unknown path
(404 when ENABLE_EXTENSION_API=false). One line per case; exit code 1 on
any failure. Finalizing writes one metadata line per case to audit.log.

Needs PyMuPDF and Pillow: run it in a throwaway client container of the app
image, for instance:
  docker run --rm --network host -e BENCH_USER -e BENCH_PASSWORD -e BENCH_INSECURE_TLS=1 \
    -e BENCH_RESOLVE_ADDRESS=192.168.1.35 -e E2E_PACE=13 -v "$PWD:/repo:ro" -w /repo \
    --entrypoint python ghcr.io/epicfail20/obfusk8-app:main benchmarks/e2e/e2e_document_flow.py
E2E_PACE (seconds before each upload/finalize) respects the existing upload
rate limit (5/min, burst 10); E2E_THEMES selects the themes (",medical").
"""

import io
import os
import re
import sys
import time
from pathlib import Path

import pymupdf as fitz
from docx import Document as WordDocument
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from obfusk8_client import BASE_URL as BASE
from obfusk8_client import login

TEXT = "Le patient Isabelle FONTAINE, née le 03/11/1980, joignable au 06 11 22 33 44, habite à Lyon."
FIXTURES = Path(__file__).resolve().parents[2] / "app" / "tests" / "fixtures"


def make_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 80), TEXT, fontsize=9)
    out = doc.tobytes()
    doc.close()
    return out


def make_png() -> bytes:
    img = Image.new("RGB", (1400, 120), "white")
    ImageDraw.Draw(img).text(
        (20, 40), "Patient Isabelle Fontaine, tel 06 11 22 33 44", fill="black", font=ImageFont.load_default(size=26)
    )
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


CASES = {
    "pdf": ("f.pdf", make_pdf),
    "docx": ("f.docx", (FIXTURES / "docx_fixture.docx").read_bytes),
    "csv": ("f.csv", (FIXTURES / "csv_fixture.csv").read_bytes),
    "image": ("f.png", make_png),
}

# Phase 2 bis: the FINAL file is checked too. Synthetic sensitive values of
# each input that must be absent from the text extracted from the downloaded
# document (the image is redacted in pixels: nothing to extract, not checked).
MUST_BE_GONE = {
    "pdf": ["FONTAINE", "06 11 22 33 44"],
    "docx": ["FONTAINE", "06 11 22 33 44", "ROUSSEAU"],
    "csv": ["marc.dupuis@example-fictif.test", "06 22 33 44 55"],
}


def final_text(kind: str, content: bytes) -> str:
    if kind == "pdf":
        with fitz.open(stream=content, filetype="pdf") as doc:
            return "".join(page.get_text() for page in doc)
    if kind == "docx":
        doc = WordDocument(io.BytesIO(content))
        cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        return "\n".join([p.text for p in doc.paragraphs] + cells)
    return content.decode("utf-8-sig")


def leaks(kind: str, content: bytes) -> str:
    """'-' for the image, else '<leaked>/<checked>' (never the values themselves)."""
    if kind not in MUST_BE_GONE:
        return "-"
    text = final_text(kind, content)
    return f"{sum(value in text for value in MUST_BE_GONE[kind])}/{len(MUST_BE_GONE[kind])}"


# upload-ratelimit (existing): 5/min, burst 10 on detect+finalize — each case
# costs 2 requests, so pace them (E2E_PACE seconds before each) when needed.
PACE = float(os.environ.get("E2E_PACE", "0"))
THEMES = os.environ.get("E2E_THEMES", ",medical").split(",")

s = login()
failures = 0
for theme in THEMES:
    for kind, (name, build) in CASES.items():
        time.sleep(PACE)
        t0 = time.monotonic()
        r = s.post(f"{BASE}/api/detect", files={"file": (name, build())}, data={"theme": theme}, timeout=120)
        dt = time.monotonic() - t0
        m = re.search(r'name="job_id" value="([0-9a-f]{32})"', r.text)
        zones = len(set(re.findall(r'data-id="([0-9a-f]{12})"', r.text)))
        if r.status_code != 200 or not m:
            failures += 1
            print(f"{kind:6} theme={theme or '-':8} detect={r.status_code} FAIL")
            continue
        if kind in ("pdf", "image"):
            p = s.get(f"{BASE}/api/preview_image/{m.group(1)}/0", timeout=60)
            preview = str(p.status_code)
        else:
            preview = "-"
        time.sleep(PACE)
        f = s.post(f"{BASE}/api/finalize", data={"job_id": m.group(1), "format": "json"}, timeout=120)
        fj = f.json()
        d = s.get(f"{BASE}{fj['download_url']}", timeout=60)
        leaked = leaks(kind, d.content) if d.status_code == 200 else "?"
        ok = f.status_code == 200 and d.status_code == 200 and len(d.content) > 0 and leaked[:2] in ("-", "0/")
        failures += not ok
        print(
            f"{kind:6} theme={theme or '-':8} detect={r.status_code} ({dt:.2f}s, {zones} zone(s)) preview={preview} "
            f"finalize={f.status_code} total={fj.get('total_redactions')} download={d.status_code} {len(d.content)}o "
            f"exposed={leaked}"
        )

ref = s.get(f"{BASE}/api/nexistepas", headers={"Accept": "application/json"}, timeout=30)
print(f"unknown path -> {ref.status_code} {ref.text}")
for method, path in (
    ("GET", "/api/v1/version"),
    ("POST", "/api/v1/text/analyze"),
    ("POST", "/api/v1/text/pseudonymize"),
):
    r = s.request(
        method, f"{BASE}{path}", json={"text": "Camille Martin"}, headers={"Accept": "application/json"}, timeout=30
    )
    print(f"{method} {path} -> {r.status_code} {r.text} x-request-id={r.headers.get('x-request-id')}")
sys.exit(1 if failures else 0)
